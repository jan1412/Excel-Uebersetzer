"""Excel lesen. Keine Neuberechnung, keine Prüfung auf veraltete Formeln."""
from __future__ import annotations

from contextlib import contextmanager
from pathlib import Path
from collections.abc import Iterator, Sequence
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .layout import Blocks, Cell, Offset
from .models import Scalar, Source, TranslationError
from .values import is_missing, normalize_value


# region E01: Arbeitsmappe mit gespeicherten Formelresultaten öffnen
@contextmanager
def open_excel(path: str | Path) -> Iterator[Any]:
    """Öffnet .xlsx/.xlsm lesend und schließt auch bei Fehlern zuverlässig."""
    path = Path(path)
    if not path.is_file():
        raise TranslationError(f"Excel-Datei nicht gefunden: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise TranslationError("Bitte eine .xlsx- oder .xlsm-Datei auswählen.")
    # Normales Laden ist bei vielen gezielten Zellzugriffen effizienter als
    # openpyxls Streaming-Modus. Es wird trotzdem nie in die Quelle geschrieben.
    workbook = load_workbook(path, data_only=True)
    try:
        yield workbook
    finally:
        workbook.close()
# endregion


# region E02: Benötigte Sheets finden
def require_sheets(workbook: Any, names: Sequence[str]) -> None:
    """Fehlende Sheets gemeinsam melden."""
    missing = sorted(set(names) - set(workbook.sheetnames))
    if missing:
        raise TranslationError("Fehlende Sheets: " + ", ".join(missing))
# endregion


# region E03: Einzelzelle mit Typ- und Pflichtwertbehandlung lesen
def read_cell(sheet: Any, spec: Cell) -> Scalar:
    """Werte/Text/Formelergebnisse lesen; echte Excel-Fehler separat melden."""
    cell = sheet[spec.address]
    source = Source(sheet.title, spec.address)
    if cell.data_type == "e":
        raise TranslationError(f"{source}: Excel-Fehler {cell.value}.")
    return normalize_value(cell.value, spec.kind, spec.required, source)
# endregion


# region E04: Relative Blockzelle in absolute Adresse übersetzen
def block_address(blocks: Blocks, origin: int, offset: Offset) -> str:
    """origin ist die tatsächliche Startzeile des aktuellen Blocks."""
    if not 0 <= offset.row < blocks.height or offset.column < 0:
        raise TranslationError(f"{blocks.sheet}: Ungültiger Offset {offset}; Blockhöhe {blocks.height}.")
    return f"{get_column_letter(blocks.first_column + offset.column)}{origin + offset.row}"
# endregion


# region E05: Relative Blockzelle lesen
def read_offset(sheet: Any, blocks: Blocks, origin: int, offset: Offset) -> Scalar:
    """Einzelwert relativ zu einem Blockanfang einlesen."""
    return read_cell(sheet, Cell(block_address(blocks, origin, offset), offset.kind, offset.required))
# endregion


# region E06: Belegte Blöcke erkennen und optionale Beschriftungen prüfen
def iter_block_rows(sheet: Any, blocks: Blocks, fields: Sequence[Offset]) -> Iterator[int]:
    """Leere Vorlagen überspringen; belegte Blöcke ohne ID ablehnen.

    Felder außerhalb last_row werden nicht eingelesen. Nur Datenfelder zählen
    für 'belegt'; statische Beschriftungen allein aktivieren keinen Block.
    """
    end = blocks.last_row if blocks.last_row is not None else sheet.max_row
    offsets = (blocks.reference, *fields)
    for origin in range(blocks.first_row, end + 1, blocks.height):
        addresses = [block_address(blocks, origin, offset) for offset in offsets]
        if not any(not is_missing(sheet[address].value) for address in addresses):
            continue
        if blocks.last_row is not None and origin + blocks.height - 1 > blocks.last_row:
            raise TranslationError(f"{blocks.sheet}: Belegter Block ab Zeile {origin} überschreitet last_row.")
        read_offset(sheet, blocks, origin, blocks.reference)
        for (row, column), expected in blocks.labels.items():
            offset = Offset(row, column, "text")
            address = block_address(blocks, origin, offset)
            actual = read_offset(sheet, blocks, origin, offset)
            if actual != expected:
                raise TranslationError(f"{sheet.title}!{address}: Beschriftung {expected!r} erwartet, erhalten: {actual!r}.")
        yield origin
# endregion


# region E07: Eine Referenz eindeutig in den passenden Namensraum aufnehmen
def add_unique(index: dict, item: Any) -> None:
    """Doppelte IDs ablehnen, statt vorherige Konfigurationen zu ersetzen."""
    if item.reference in index:
        previous = index[item.reference]
        raise TranslationError(f"{item.source}: Doppelte Referenz {item.reference!r}; zuerst bei {previous.source}.")
    index[item.reference] = item
# endregion
