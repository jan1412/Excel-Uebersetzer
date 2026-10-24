"""Ausführen: python -m unittest discover -s tests -v

Die Test-Arbeitsmappe existiert nur im Speicher; es wird keine Excel-Datei
erzeugt oder neu berechnet. Der Dateiöffner wird mit genau dieser Mappe geprüft.
"""
from dataclasses import replace
from decimal import Decimal as D
import ast
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import Workbook

import main as application
from examples.cfg_demo import build_demo_project
from excel_uebersetzer.dropdowns import read_dropdown_id, selection_to_id
from excel_uebersetzer.excel_reader import open_excel
from excel_uebersetzer.layout import MatrixSlot, Offset
from excel_uebersetzer.loaders import read_project
from excel_uebersetzer.mapping import LAYOUT
from excel_uebersetzer.models import MatrixReference, Source, TranslationError
from excel_uebersetzer.pipeline import translate
from excel_uebersetzer.renderer import document_chunks, python_literal, write_document
from excel_uebersetzer.resolver import resolve_steps, select_matrix, validate_project
from excel_uebersetzer.values import format_number, is_missing, normalize_value, positive_index, split_references, to_decimal

ROOT = Path(__file__).resolve().parent.parent


def example_workbook():
    """Repräsentative Excel-Zelltypen und Blockanordnung, ohne Festplattendatei."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, cells in json.loads((ROOT / "examples/cfg_cells.json").read_text(encoding="utf-8")).items():
        sheet = workbook.create_sheet(sheet_name)
        for address, value in cells.items():
            sheet[address] = value
    return workbook


class ValueTests(unittest.TestCase):
    def test_missing_values(self):
        for value in (None, "", "   ", "ka", " KA "):
            with self.subTest(value=value):
                self.assertTrue(is_missing(value))
                self.assertIsNone(normalize_value(value))

    def test_zero_and_false_are_not_missing(self):
        for value in (0, 0.0, D("0"), False, "0"):
            with self.subTest(value=value):
                self.assertFalse(is_missing(value))
        self.assertEqual(normalize_value(0), D("0"))
        self.assertIs(normalize_value(False), False)

    def test_decimal_point_and_precision(self):
        self.assertEqual(to_decimal("1.234567890123456789"), D("1.234567890123456789"))
        self.assertEqual(to_decimal(0.1), D("0.1"))
        self.assertEqual(to_decimal("1e-3"), D("0.001"))

    def test_invalid_numbers(self):
        for value in ("1,25", "abc", True, float("inf"), float("nan")):
            with self.subTest(value=value), self.assertRaises(TranslationError):
                to_decimal(value)

    def test_types_and_required(self):
        self.assertEqual(normalize_value("001", "text"), "001")
        self.assertEqual(normalize_value("1.00", "auto"), "1.00")
        self.assertEqual(normalize_value("1.00", "number"), D("1.00"))
        self.assertIs(normalize_value("Nein", "bool"), False)
        with self.assertRaises(TranslationError):
            normalize_value("ka", "number", True)

    def test_number_output(self):
        for value, expected in ((D("0"), "0.00"), (D("1.23456789"), "1.23456789"),
                                (D("1E+3"), "1000.00"), (D("1E-7"), "0.0000001")):
            with self.subTest(value=value):
                self.assertEqual(format_number(value), expected)

    def test_indices_do_not_round(self):
        self.assertEqual(positive_index(D("1.00"), "Test"), 1)
        for value in (D("1.50"), D("0")):
            with self.assertRaises(TranslationError):
                positive_index(value, "Test")

    def test_reference_lists(self):
        self.assertEqual(split_references("CFG001; CFG002", "Test"), ("CFG001", "CFG002"))
        self.assertEqual(split_references(None, "Test"), ())
        with self.assertRaises(TranslationError):
            split_references("CFG001;CFG001", "Test")


class ReaderTests(unittest.TestCase):
    def setUp(self):
        self.workbook = example_workbook()

    def tearDown(self):
        self.workbook.close()

    def test_five_sheets_and_repeated_blocks(self):
        project = read_project(self.workbook, LAYOUT)
        self.assertEqual((len(project.materials), len(project.settings), len(project.machines), len(project.steps)), (2, 2, 1, 2))
        self.assertEqual(project.materials["MAT_STEEL_001"].material, project.materials["MAT_STEEL_002"].material)
        self.assertEqual(project.settings["CFG001"].parameters["limits"], (D("0"), D("1")))
        self.assertEqual(project.steps[0].parameters["feed"], D("0"))
        self.assertIsNone(project.steps[1].parameters["feed"])

    def test_demo_and_excel_readers_agree(self):
        project = read_project(self.workbook, LAYOUT)
        expected = build_demo_project()
        from excel_uebersetzer.renderer import step_payload
        self.assertEqual(project.metadata, expected.metadata)
        self.assertEqual([step_payload(s) for s in resolve_steps(project)],
                         [step_payload(s) for s in resolve_steps(expected)])

    def test_missing_sheet(self):
        self.workbook.remove(self.workbook["Metadata"])
        with self.assertRaisesRegex(TranslationError, "Metadata"):
            read_project(self.workbook, LAYOUT)

    def test_duplicate_material_reference(self):
        self.workbook["Materialien"]["A6"] = "MAT_STEEL_001"
        with self.assertRaisesRegex(TranslationError, "Doppelte Referenz"):
            read_project(self.workbook, LAYOUT)

    def test_duplicate_step_reference(self):
        self.workbook["Ablauferstellung"]["A8"] = "STEP001"
        with self.assertRaisesRegex(TranslationError, "Doppelte Referenz"):
            read_project(self.workbook, LAYOUT)

    def test_partial_block_without_id(self):
        self.workbook["Materialien"]["A6"] = None
        with self.assertRaisesRegex(TranslationError, "Materialien!A6"):
            read_project(self.workbook, LAYOUT)

    def test_blank_template_is_ignored(self):
        self.workbook["Materialien"]["A11"] = "Dichte"
        self.workbook["Materialien"]["A12"] = "Dicke"
        self.assertEqual(len(read_project(self.workbook, LAYOUT).materials), 2)

    def test_copy_paste_adds_third_configuration(self):
        sheet = self.workbook["Einstellungen"]
        for row in range(2, 6):
            for column in range(1, 4):
                sheet.cell(row + 8, column).value = sheet.cell(row, column).value
        sheet["A10"] = "CFG003"
        self.workbook["Ablauferstellung"]["B10"] = "CFG003"
        project = read_project(self.workbook, LAYOUT)
        self.assertEqual(len(project.settings), 3)
        self.assertEqual(resolve_steps(project)[1].settings[0].reference, "CFG003")

    def test_optional_slots_keep_position(self):
        self.workbook["Einstellungen"]["B5"] = "ka"
        self.assertEqual(read_project(self.workbook, LAYOUT).settings["CFG001"].parameters["limits"], (None, D("1")))

    def test_partial_matrix_bounds(self):
        self.workbook["Ablauferstellung"]["E6"] = None
        with self.assertRaisesRegex(TranslationError, "Alle vier"):
            read_project(self.workbook, LAYOUT)

    def test_bounds_without_matrix_id(self):
        self.workbook["Ablauferstellung"]["B5"] = None
        with self.assertRaisesRegex(TranslationError, "ohne Konfigurations-ID"):
            read_project(self.workbook, LAYOUT)

    def test_multiple_matrix_slots(self):
        self.workbook["Ablauferstellung"]["C5"] = "MCH001"
        step_layout = replace(LAYOUT.steps, matrix_slots=(*LAYOUT.steps.matrix_slots, MatrixSlot(Offset(3, 2, "text"))))
        project = read_project(self.workbook, replace(LAYOUT, steps=step_layout))
        matrices = resolve_steps(project)[0].matrices
        self.assertEqual([len(matrix.values) for matrix in matrices], [2, 3])

    def test_excel_error_is_not_missing(self):
        self.workbook["Einstellungen"]["B3"] = "#DIV/0!"
        with self.assertRaisesRegex(TranslationError, "Excel-Fehler"):
            read_project(self.workbook, LAYOUT)

    def test_optional_label_check(self):
        material_layout = replace(LAYOUT.materials, blocks=replace(LAYOUT.materials.blocks, labels={(1, 0): "Dichte"}))
        self.workbook["Materialien"]["A7"] = "Falsch kopiert"
        with self.assertRaisesRegex(TranslationError, "Beschriftung"):
            read_project(self.workbook, replace(LAYOUT, materials=material_layout))


class ResolverTests(unittest.TestCase):
    def test_matrix_selection(self):
        project = build_demo_project()
        resolved = resolve_steps(project)
        self.assertEqual(resolved[0].matrices[0].values, ((D("2"), D("3")), (D("5"), D("6"))))
        self.assertEqual(resolved[1].matrices[0].values, project.machines["MCH001"].matrix)

    def test_out_of_bounds(self):
        project = build_demo_project()
        with self.assertRaises(TranslationError):
            select_matrix(project.machines["MCH001"].matrix, MatrixReference("MCH001", (1, 1, 4, 2)))

    def test_unknown_reference(self):
        project = build_demo_project()
        project.steps[0].setting_references = ("CFG999",)
        with self.assertRaisesRegex(TranslationError, "CFG999"):
            resolve_steps(project)

    def test_empty_steps_allowed(self):
        project = build_demo_project()
        project.steps = []
        validate_project(project, resolve_steps(project))
        output = "".join(document_chunks(project.metadata, [], "py"))
        compile(output, "demo.py", "exec")


class DropdownTests(unittest.TestCase):
    def setUp(self):
        self.workbook = example_workbook()

    def tearDown(self):
        self.workbook.close()

    def test_normal_dropdown_id(self):
        self.assertEqual(read_dropdown_id(self.workbook["Ablauferstellung"], "B4"), "CFG001")

    def test_form_control_index(self):
        self.workbook["Ablauferstellung"]["G2"] = 2.0
        self.workbook["Einstellungen"]["H2"] = "CFG001"
        self.workbook["Einstellungen"]["H3"] = "CFG002"
        self.assertEqual(read_dropdown_id(self.workbook["Ablauferstellung"], "G2",
                                         self.workbook["Einstellungen"], "H2:H3"), "CFG002")

    def test_blank_options_keep_their_index(self):
        self.assertEqual(selection_to_id(D("3"), ("CFG001", None, "CFG003")), "CFG003")
        with self.assertRaises(TranslationError):
            selection_to_id(D("2"), ("CFG001", None, "CFG003"))

    def test_index_outside_options(self):
        with self.assertRaises(TranslationError):
            selection_to_id(D("4"), ("CFG001", "CFG002"))

    def test_single_cell_source(self):
        self.workbook["Ablauferstellung"]["G2"] = 1.0
        self.assertEqual(read_dropdown_id(self.workbook["Ablauferstellung"], "G2",
                                         self.workbook["Einstellungen"], "A2"), "CFG001")

    def test_multiple_source_columns_rejected(self):
        with self.assertRaises(TranslationError):
            read_dropdown_id(self.workbook["Ablauferstellung"], "B4", self.workbook["Einstellungen"], "A2:B2")


class OutputTests(unittest.TestCase):
    def test_text_output_sequence_zero_and_missing(self):
        project = build_demo_project()
        text = "".join(document_chunks(project.metadata, resolve_steps(project)))
        self.assertLess(text.index("Step 1: STEP001"), text.index("Step 2: STEP002"))
        self.assertIn("feed = 0.00", text)
        self.assertIn("feed = ka", text)
        self.assertIn("1500.00", text)

    def test_python_output_is_valid_and_text_stays_data(self):
        project = build_demo_project()
        dangerous = "Zitat ' \\\n__import__('os').system('echo NIE AUSFÜHREN')"
        project.steps[0].name = dangerous
        output = "".join(document_chunks(project.metadata, resolve_steps(project), "py"))
        compile(output, "demo.py", "exec")
        tree = ast.parse(output)
        assignment = next(node for node in tree.body if isinstance(node, ast.Assign) and node.targets[0].id == "STEPS")
        values = ast.literal_eval(assignment.value)
        self.assertEqual(values[0]["name"], dangerous)
        self.assertEqual(values[0]["parameters"]["feed"], 0.0)
        self.assertIsNone(values[1]["parameters"]["feed"])

    def test_high_precision_literal_is_not_rounded(self):
        self.assertEqual(python_literal(D("1.123456789123456789")), "1.123456789123456789")

    def test_template_failure_preserves_existing_file(self):
        with TemporaryDirectory() as temp:
            target = Path(temp) / "output.txt"
            target.write_text("Vorhanden", encoding="utf-8")
            def broken_chunks():
                yield "Anfang"
                raise TranslationError("Vorlagenfehler")
            with self.assertRaises(TranslationError):
                write_document(target, broken_chunks())
            self.assertEqual(target.read_text(encoding="utf-8"), "Vorhanden")
            self.assertEqual(list(Path(temp).iterdir()), [target])

    def test_invalid_python_preserves_existing_file(self):
        with TemporaryDirectory() as temp:
            target = Path(temp) / "output.py"
            target.write_text("x = 1\n", encoding="utf-8")
            with self.assertRaisesRegex(TranslationError, "Python-Ausgabe ungültig"):
                write_document(target, iter(["def :"]), "py")
            self.assertEqual(target.read_text(), "x = 1\n")

    def test_missing_placeholder_is_reported(self):
        from excel_uebersetzer.renderer import render_template
        with TemporaryDirectory() as temp:
            template = Path(temp) / "step.tpl"
            template.write_text("${not_defined}", encoding="utf-8")
            with self.assertRaisesRegex(TranslationError, "not_defined"):
                render_template(template, {})

    def test_pipeline_reads_saved_values_and_does_not_write_excel(self):
        with TemporaryDirectory() as temp:
            source = Path(temp) / "input.xlsx"
            source.write_bytes(b"Nur ein Platzhalter; Reader wird isoliert getestet.")
            output = Path(temp) / "result.py"
            with patch("excel_uebersetzer.excel_reader.load_workbook", return_value=example_workbook()) as loader:
                translate(source, output, "py")
                loader.assert_called_once_with(source, data_only=True)
            compile(output.read_text(encoding="utf-8"), str(output), "exec")
            self.assertEqual(source.read_bytes(), b"Nur ein Platzhalter; Reader wird isoliert getestet.")

    def test_input_cannot_be_output(self):
        with self.assertRaisesRegex(TranslationError, "nicht die Excel"):
            translate("input.xlsx", "input.xlsx")

    def test_open_excel_closes_on_error(self):
        with TemporaryDirectory() as temp:
            source = Path(temp) / "input.xlsx"
            source.touch()
            workbook = example_workbook()
            with patch("excel_uebersetzer.excel_reader.load_workbook", return_value=workbook), patch.object(workbook, "close") as close:
                with self.assertRaisesRegex(RuntimeError, "absichtlich"):
                    with open_excel(source):
                        raise RuntimeError("absichtlich")
                close.assert_called_once()


class ApplicationTests(unittest.TestCase):
    def test_cancel_file_picker(self):
        with patch("main.choose_excel", return_value=None), patch("main.translate") as convert:
            self.assertEqual(application.main([]), 0)
            convert.assert_not_called()

    def test_cli_python_format_detection(self):
        with patch("main.translate", return_value=Path("out.py")) as convert, patch("main.show_message") as message:
            self.assertEqual(application.main(["input.xlsx", "--output", "out.py"]), 0)
            self.assertEqual(convert.call_args.args[2], "py")
            message.assert_not_called()

    def test_save_dialog_python_format_detection(self):
        with patch("main.choose_excel", return_value=Path("input.xlsx")), patch("main.choose_output", return_value=Path("out.py")), patch("main.translate", return_value=Path("out.py")) as convert, patch("main.show_message"):
            self.assertEqual(application.main([]), 0)
            self.assertEqual(convert.call_args.args[2], "py")

    def test_explicit_format_wins_over_extension(self):
        with patch("main.translate", return_value=Path("out.data")) as convert:
            self.assertEqual(application.main(["input.xlsx", "--output", "out.data", "--format", "py"]), 0)
            self.assertEqual(convert.call_args.args[2], "py")


if __name__ == "__main__":
    unittest.main()
